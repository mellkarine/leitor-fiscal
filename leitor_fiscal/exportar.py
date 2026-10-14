"""Exportação para Excel com formatação (cabeçalho roxo, datas e valores de verdade)."""
import io

import pandas as pd
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

ROXO = "6D28D9"
COLUNAS_DOCS = {
    "arquivo": "Arquivo", "tipo": "Tipo", "emitente": "Emitente", "doc_emitente": "CNPJ/CPF do emitente",
    "numero": "Número", "data": "Data", "vencimento": "Vencimento", "valor_total": "Valor total",
    "destinatario": "Destinatário/pagador", "linha_digitavel": "Linha digitável", "status": "Status",
    "alertas": "Alertas",
}
COLUNAS_ITENS = {
    "arquivo": "Arquivo", "descricao": "Descrição", "quantidade": "Quantidade",
    "valor_unitario": "Valor unitário", "valor_item": "Valor do item",
}
STATUS = {"ok": "OK", "revisar": "Revisar", "erro": "Erro"}
COR_STATUS = {"OK": "E8F7EE", "Revisar": "FFF4DE", "Erro": "FDE8E8"}
FORMATO_MOEDA = '"R$" #,##0.00'
FORMATO_DATA = "DD/MM/YYYY"


def _preparar(df: pd.DataFrame, colunas: dict[str, str]) -> pd.DataFrame:
    """Reordena e renomeia as colunas presentes; converte datas e tipos para o Excel."""
    df = df[[c for c in colunas if c in df.columns]].copy()
    for c in ("data", "vencimento"):
        if c in df.columns:
            df[c] = pd.to_datetime(df[c], errors="coerce").dt.date
    if "tipo" in df.columns:
        df["tipo"] = df["tipo"].astype("string").str.replace("_", " ")
    if "status" in df.columns:
        df["status"] = df["status"].map(STATUS).fillna(df["status"])
    return df.rename(columns=colunas)


def _formatar(ws, moeda: set[str], datas: set[str]) -> None:
    cabecalho = {c.value: c.column for c in ws[1]}
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor=ROXO)
        c.alignment = Alignment(vertical="center")
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    ws.row_dimensions[1].height = 22
    for nome, col in cabecalho.items():
        letra = get_column_letter(col)
        largura = max(len(str(c.value or "")) for c in ws[letra])
        ws.column_dimensions[letra].width = min(max(largura + 3, 12), 60)
        for celula in ws[letra][1:]:
            if nome in moeda:
                celula.number_format = FORMATO_MOEDA
            elif nome in datas:
                celula.number_format = FORMATO_DATA
            elif nome == "Status" and celula.value in COR_STATUS:
                celula.fill = PatternFill("solid", fgColor=COR_STATUS[celula.value])
                celula.font = Font(bold=True)
            elif nome == "Alertas":
                celula.alignment = Alignment(wrap_text=True, vertical="top")


def gerar_excel(docs: pd.DataFrame, itens: pd.DataFrame) -> bytes:
    """Excel com as abas 'documentos', 'itens' e 'resumo'."""
    docs_x = _preparar(docs, COLUNAS_DOCS)
    itens_x = _preparar(itens, COLUNAS_ITENS) if not itens.empty else pd.DataFrame(columns=list(COLUNAS_ITENS.values()))

    status = docs["status"] if "status" in docs.columns else pd.Series(dtype=str)
    valores = pd.to_numeric(docs["valor_total"], errors="coerce") if "valor_total" in docs.columns else pd.Series(dtype=float)
    resumo = pd.DataFrame({
        "Indicador": ["Documentos lidos", "Sem alertas", "Para revisão", "Erros", "Valor total extraído"],
        "Valor": [len(docs), int((status == "ok").sum()), int((status == "revisar").sum()),
                  int((status == "erro").sum()), float(valores.sum())],
    })

    buffer = io.BytesIO()
    with pd.ExcelWriter(buffer, engine="openpyxl") as xl:
        docs_x.to_excel(xl, sheet_name="documentos", index=False)
        itens_x.to_excel(xl, sheet_name="itens", index=False)
        resumo.to_excel(xl, sheet_name="resumo", index=False)
        _formatar(xl.sheets["documentos"], {"Valor total"}, {"Data", "Vencimento"})
        _formatar(xl.sheets["itens"], {"Valor unitário", "Valor do item"}, set())
        _formatar(xl.sheets["resumo"], set(), set())
        xl.sheets["resumo"]["B6"].number_format = FORMATO_MOEDA
    return buffer.getvalue()
