"""Testes que rodam sem chamar a API: texto dos PDFs, schema, validações, demo, CLI e avaliação."""
import io
from datetime import date
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook

from leitor_fiscal.avaliar import avaliar
from leitor_fiscal.cli import processar
from leitor_fiscal.demo import cadeia_demo
from leitor_fiscal.exportar import gerar_excel
from leitor_fiscal.pdf_texto import extrair_texto
from leitor_fiscal.pipeline import para_dataframes, processar_lote
from leitor_fiscal.schemas import Documento, Item
from leitor_fiscal.validacao import cnpj_valido, conferencias, cpf_valido, formatar_brl, validar

RAIZ = Path(__file__).parent.parent
PDFS = RAIZ / "data" / "sintetico" / "pdfs"
GABARITO = RAIZ / "data" / "sintetico" / "gabarito" / "documentos.csv"


def _lote():
    return processar_lote([(p.name, p) for p in sorted(PDFS.glob("*.pdf"))], cadeia_demo())


def test_texto_dos_pdfs_e_selecionavel():
    assert "NOTA FISCAL DE VENDA" in extrair_texto(PDFS / "nota_001.pdf")
    assert "Linha digitável" in extrair_texto(PDFS / "boleto_001.pdf")
    assert "RECIBO" in extrair_texto(PDFS / "recibo_001.pdf")


def test_schema_aceita_documento_completo_e_campos_vazios():
    nota = Documento(tipo="nota_fiscal", data="2026-09-14", valor_total=448.0,
                     itens=[Item(descricao="Grampeador", quantidade=2, valor_unitario=34.5, valor_item=69.0)])
    assert nota.data == date(2026, 9, 14)
    assert nota.vencimento is None  # ausente fica vazio, nunca inventado


def test_cnpj_e_cpf():
    assert cnpj_valido("45.012.367/0001-63")
    assert not cnpj_valido("45.012.367/0001-64")
    assert not cnpj_valido("11.111.111/1111-11")
    assert cpf_valido("529.982.247-25")
    assert not cpf_valido("529.982.247-26")


def test_formatar_brl():
    assert formatar_brl(1234.5) == "R$ 1.234,50"
    assert formatar_brl(448) == "R$ 448,00"


def test_validacoes_dos_casos_dificeis():
    r = {x.arquivo: x for x in _lote()}
    assert r["nota_001.pdf"].alertas == [] and r["boleto_001.pdf"].alertas == [] and r["recibo_001.pdf"].alertas == []
    assert any("diverge da soma dos itens" in a for a in r["nota_002.pdf"].alertas)
    assert any("CNPJ do emitente inválido" in a for a in r["nota_003.pdf"].alertas)
    assert any("vencimento" in a for a in r["boleto_002.pdf"].alertas)
    assert any("valor total" in a for a in r["recibo_002.pdf"].alertas)


def test_cli_em_modo_demo(tmp_path):
    resumo = processar(PDFS, tmp_path, cadeia=cadeia_demo())
    assert resumo == {"pdfs": 7, "ok": 3, "revisar": 4, "erro": 0}
    docs = pd.read_csv(tmp_path / "documentos.csv")
    assert set(docs["tipo"]) == {"nota_fiscal", "boleto", "recibo"}
    assert len(pd.read_csv(tmp_path / "itens.csv")) == 7


def test_pdf_invalido_nao_derruba_o_lote():
    registros = processar_lote([("quebrado.pdf", b"isto nao e um pdf"), ("nota_001.pdf", PDFS / "nota_001.pdf")], cadeia_demo())
    assert registros[0].status == "erro"
    assert registros[1].status == "ok"


def test_avaliar_mede_acerto_contra_o_gabarito(tmp_path):
    processar(PDFS, tmp_path, cadeia=cadeia_demo())
    previsto = pd.read_csv(tmp_path / "documentos.csv", dtype=str)
    gabarito = pd.read_csv(GABARITO, dtype=str)
    assert avaliar(previsto, gabarito).iloc[-1]["acerto_%"] == 100.0
    previsto.loc[previsto["arquivo"] == "nota_001.pdf", "valor_total"] = "999.00"
    assert avaliar(previsto, gabarito).iloc[-1]["acerto_%"] < 100.0


def test_conferencias_mostram_o_que_passou_e_o_que_falhou():
    r = {x.arquivo: x for x in _lote()}
    ok = {c.rotulo: c.ok for c in conferencias(r["nota_001.pdf"].documento)}
    assert ok == {"Campos obrigatórios": True, "CNPJ/CPF do emitente": True, "Itens e total": True}
    ruim = {c.rotulo: c.ok for c in conferencias(r["nota_002.pdf"].documento)}
    assert ruim["Itens e total"] is False and ruim["CNPJ/CPF do emitente"] is True
    assert "Datas" in {c.rotulo for c in conferencias(r["boleto_001.pdf"].documento)}  # só boleto tem a verificação


def test_validar_e_a_soma_das_conferencias():
    for reg in _lote():
        assert reg.alertas == [a for c in conferencias(reg.documento) for a in c.alertas]


def test_item_sem_subtotal_usa_quantidade_vezes_unitario():
    doc = Documento(tipo="nota_fiscal", emitente="Aurora Ltda", doc_emitente="45.012.367/0001-63", numero="1",
                    data="2026-09-14", valor_total=69.0,
                    itens=[Item(descricao="Grampeador", quantidade=2, valor_unitario=34.5)])
    assert validar(doc) == []


def test_item_sem_nenhum_valor_gera_alerta_claro():
    doc = Documento(tipo="nota_fiscal", emitente="Aurora Ltda", doc_emitente="45.012.367/0001-63", numero="1",
                    data="2026-09-14", valor_total=69.0, itens=[Item(descricao="Grampeador")])
    assert validar(doc) == ["Item 'Grampeador' sem valor"]  # sem falso 'total diverge'


def test_arquivos_com_o_mesmo_nome_nao_se_sobrescrevem():
    pdf = PDFS / "nota_001.pdf"
    registros = processar_lote([("nota.pdf", pdf), ("nota.pdf", pdf), ("nota.pdf", pdf)], cadeia_demo())
    assert [r.arquivo for r in registros] == ["nota.pdf", "nota (2).pdf", "nota (3).pdf"]
    assert all(r.status == "ok" for r in registros)


def test_registro_guarda_o_texto_lido():
    registros = processar_lote([("nota_001.pdf", PDFS / "nota_001.pdf")], cadeia_demo())
    assert "NOTA FISCAL DE VENDA" in registros[0].texto


def test_excel_formatado_tem_tres_abas_e_valores_numericos():
    docs, itens = para_dataframes(_lote())
    wb = load_workbook(io.BytesIO(gerar_excel(docs, itens)))
    assert wb.sheetnames == ["documentos", "itens", "resumo"]
    ws = wb["documentos"]
    cab = {c.value: c.column for c in ws[1]}
    assert ws.cell(2, cab["Valor total"]).number_format == '"R$" #,##0.00'
    assert isinstance(ws.cell(2, cab["Valor total"]).value, float)
    assert ws.cell(2, cab["Data"]).number_format == "DD/MM/YYYY"
    assert {ws.cell(i, cab["Status"]).value for i in range(2, ws.max_row + 1)} == {"OK", "Revisar"}


def test_excel_aceita_lote_so_com_erros():
    docs, itens = para_dataframes(processar_lote([("quebrado.pdf", b"isto nao e um pdf")], cadeia_demo()))
    wb = load_workbook(io.BytesIO(gerar_excel(docs, itens)))
    assert wb["resumo"]["B5"].value == 1  # linha "Erros"
